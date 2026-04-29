"""
SAPBA Docx Exporter Tool Injector
Injects a custom Python Tool into Open WebUI for generating Word documents.
"""
import requests, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

BASE_URL = 'http://localhost:3013'
TOKEN = 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpZCI6IjNiNDhhMWQ5LTc3NzAtNDQ0Ni05Mzg4LTQxOWU2ZDY5YzEwYSIsImV4cCI6MTc3OTQ5MTIzNywianRpIjoiN2Y3NjhlNmUtMjA0OC00NjMwLTk5NzgtMzM2YzU2ZWE5NThkIiwiaWF0IjoxNzc3MDcyMDM3fQ.eJ7VE_jfFilMPNwefsZNiPdM_drjnSCN5UgDxLFCoGg'
HEADERS = {'Authorization': f'Bearer {TOKEN}', 'Content-Type': 'application/json'}

TOOL_CONTENT = '''"""
title: SAPBA Word Exporter
author: Antigravity Architect
description: Automatically generates and provides a download link for a Word document (.docx) based on the provided markdown text.
version: 1.0.0
"""

class Tools:
    def __init__(self):
        pass

    def export_to_docx(self, markdown_text: str, title: str) -> str:
        """
        Exports the provided markdown text to a Word Document (.docx) and returns a clickable HTML download link.
        Call this tool ONLY when the user explicitly asks to generate or download a document/file (e.g. Word, Docx).
        
        :param markdown_text: The complete text content to export. Ensure this is the full and final text.
        :param title: The title of the document (used for the filename).
        """
        import io
        from docx import Document
        import base64
        
        try:
            doc = Document()
            doc.add_heading(title, 0)
            
            for line in markdown_text.split('\\n'):
                line = line.strip()
                if not line:
                    continue
                if line.startswith('# '):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], level=3)
                elif line.startswith('- '):
                    doc.add_paragraph(line[2:], style='List Bullet')
                elif line.startswith('1. ') or line.startswith('2. '):
                    doc.add_paragraph(line[3:], style='List Number')
                else:
                    doc.add_paragraph(line)
            
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            
            b64 = base64.b64encode(buffer.read()).decode('utf-8')
            mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            safe_title = title.replace(" ", "_").replace("/", "_")
            html_link = f'<a href="data:{mime};base64,{b64}" download="{safe_title}.docx" style="display: inline-block; padding: 10px 20px; background-color: #2563eb; color: white; text-decoration: none; border-radius: 5px; font-weight: bold; margin-top: 10px;">📄 Download {safe_title}.docx</a>'
            
            return f"Dokumen berhasil dibuat! Silakan klik tombol di bawah ini untuk mengunduhnya:\\n\\n{html_link}"
        except Exception as e:
            return f"Terjadi kesalahan saat membuat dokumen: {str(e)}"
'''

payload = {
    "id": "sapba_docx_exporter",
    "name": "SAPBA Word Exporter",
    "meta": {
        "description": "Ekspor naskah ke Microsoft Word (.docx)",
    },
    "content": TOOL_CONTENT
}

r = requests.get(f'{BASE_URL}/api/v1/tools/id/sapba_docx_exporter', headers=HEADERS)
if r.ok:
    print("Tool exists. Updating...")
    r2 = requests.post(f'{BASE_URL}/api/v1/tools/id/sapba_docx_exporter/update', headers=HEADERS, json=payload)
    if r2.ok:
        print("[OK] Tool updated successfully.")
    else:
        print(f"[FAIL] Update failed: {r2.text}")
else:
    print("Creating new tool...")
    r2 = requests.post(f'{BASE_URL}/api/v1/tools/create', headers=HEADERS, json=payload)
    if r2.ok:
        print("[OK] Tool created successfully.")
    else:
        print(f"[FAIL] Create failed: {r2.text}")

